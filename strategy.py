import io
from datetime import datetime, timezone

import numpy as np
import pandas as pd

from signals import rsi14, atr14

PCT_METRICS = {"總回報", "年化回報 (CAGR)", "年化波動率", "最大回撤", "勝率（已平倉）", "市場參與度"}
RULES = [
    ("入場條件（全部成立）", "調整收市價 > 50 日均線；50 日均線 > 200 日均線；MACD 柱狀圖 > 0；RSI14 介乎 40 至 70；ATR 有效"),
    ("出場條件（任一成立）", "調整收市價跌穿 50 日均線；或調整收市價觸及追蹤止損（入場後最高的 收市價 − 止損倍數 × ATR）"),
    ("執行假設", "訊號在收市後確認，下一個交易日收市價成交；每次買入或賣出各扣一次交易成本；只做多、不使用槓桿"),
    ("風險回報比", "(目標價 − 參考入場價) ÷ (參考入場價 − 止損價)；目標價 = 收市價 + 目標倍數 × ATR；另列 60 日高位作阻力參考"),
    ("倉位計算", "建議股數 = 帳戶資金 × 單筆風險% ÷ (參考入場價 − 止損價)，向下取整，未考慮資金上限與最小交易單位"),
    ("限制", "規則為簡單技術系統，歷史績效不代表未來；未包含基本面、估值、財報日、股息稅、匯率、滑點與流動性影響"),
]


def add_indicators(df):
    d = df.copy()
    p = d["Adj Close"]
    ema12, ema26 = p.ewm(span=12, adjust=False).mean(), p.ewm(span=26, adjust=False).mean()
    d["MACD"] = ema12 - ema26
    d["MACD_Signal"] = d["MACD"].ewm(span=9, adjust=False).mean()
    d["MACD_Hist"] = d["MACD"] - d["MACD_Signal"]
    d["RSI14"] = rsi14(p)
    d["ATR14"] = atr14(d)
    d["ATR_Adj"] = d["ATR14"] / d["Close"] * p
    mid = p.rolling(20, min_periods=20).mean()
    sd = p.rolling(20, min_periods=20).std(ddof=0)
    d["BB_Mid"], d["BB_Upper"], d["BB_Lower"] = mid, mid + 2 * sd, mid - 2 * sd
    d["High60"] = d["High"].rolling(60, min_periods=60).max()
    return d


def generate_signals(d, stop_mult):
    p, atr = d["Adj Close"].to_numpy(), d["ATR_Adj"].to_numpy()
    entry = ((d["Adj Close"] > d["SMA_50_Adj"]) & (d["SMA_50_Adj"] > d["SMA_200_Adj"]) &
             (d["MACD_Hist"] > 0) & d["RSI14"].between(40, 70) & d["ATR_Adj"].notna()).to_numpy()
    below50 = (d["Adj Close"] < d["SMA_50_Adj"]).to_numpy()
    n = len(d)
    sig, reason = np.zeros(n, dtype=int), [""] * n
    in_pos, stop = False, np.nan
    for i in range(n):
        if not in_pos:
            if entry[i]:
                in_pos, stop = True, p[i] - stop_mult * atr[i]
                sig[i], reason[i] = 1, "入場"
        else:
            if np.isfinite(atr[i]):
                stop = max(stop, p[i] - stop_mult * atr[i])
            if p[i] <= stop:
                in_pos, reason[i] = False, "出場：觸及追蹤止損"
            elif below50[i]:
                in_pos, reason[i] = False, "出場：跌穿 50 日均線"
            else:
                sig[i] = 1
    return pd.DataFrame({"Entry_Cond": entry, "Signal": sig, "Event": reason}, index=d.index)


def perf(r):
    r = r.fillna(0.0)
    eq = (1 + r).cumprod()
    years = len(r) / 252
    sd = r.std(ddof=1)
    return {"總回報": eq.iloc[-1] - 1,
            "年化回報 (CAGR)": eq.iloc[-1] ** (1 / years) - 1 if years > 0 and eq.iloc[-1] > 0 else np.nan,
            "年化波動率": sd * np.sqrt(252),
            "夏普比率（無風險利率 0）": r.mean() / sd * np.sqrt(252) if sd > 0 else np.nan,
            "最大回撤": (eq / eq.cummax() - 1).min()}


def backtest(d, signal, cost_bps):
    p = d["Adj Close"]
    ret = p.pct_change().fillna(0.0)
    pos = signal.shift(2).fillna(0).astype(int)
    turnover = pos.diff().abs().fillna(pos.abs())
    cost = cost_bps / 10000
    sret = pos * ret - turnover * cost
    daily = pd.DataFrame({"Position": pos, "Strategy_Return": sret,
                          "Strategy_Equity": (1 + sret).cumprod(), "BuyHold_Equity": (1 + ret).cumprod()})
    pv, trades, n = pos.to_numpy(), [], len(pos)
    start = None
    for i in range(n):
        if pv[i] == 1 and (i == 0 or pv[i - 1] == 0):
            start = i
        if pv[i] == 1 and (i == n - 1 or pv[i + 1] == 0) and start is not None:
            e = max(start - 1, 0)
            gross = p.iloc[i] / p.iloc[e] - 1
            trades.append({"入場日": p.index[e].date(), "入場價（調整）": p.iloc[e], "出場日": p.index[i].date(),
                           "出場價（調整）": p.iloc[i], "持有交易日": i - start + 1,
                           "扣成本後回報": (1 + gross) * (1 - cost) ** 2 - 1,
                           "狀態": "持倉中" if i == n - 1 else "已平倉"})
            start = None
    tr = pd.DataFrame(trades, columns=["入場日", "入場價（調整）", "出場日", "出場價（調整）",
                                      "持有交易日", "扣成本後回報", "狀態"])
    closed = tr[tr["狀態"] == "已平倉"]["扣成本後回報"]
    wins, losses = closed[closed > 0], closed[closed <= 0]
    sp, bp = perf(sret), perf(ret)
    rows = []
    for k in sp:
        rows.append((k, sp[k], bp[k]))
    rows += [("交易次數（含持倉中）", len(tr), np.nan),
             ("勝率（已平倉）", len(wins) / len(closed) if len(closed) else np.nan, np.nan),
             ("平均盈利／平均虧損", wins.mean() / abs(losses.mean()) if len(wins) and len(losses) and losses.mean() != 0 else np.nan, np.nan),
             ("盈虧比 Profit Factor", wins.sum() / abs(losses.sum()) if len(losses) and losses.sum() != 0 else np.nan, np.nan),
             ("市場參與度", float(pos.mean()), 1.0)]
    return daily, tr, pd.DataFrame(rows, columns=["指標", "策略", "買入持有"])


def recommend(d, sg, params):
    last, prev = d.iloc[-1], d.iloc[-2] if len(d) > 1 else d.iloc[-1]
    s_now, s_prev = int(sg["Signal"].iloc[-1]), int(sg["Signal"].iloc[-2]) if len(sg) > 1 else 0
    conds = [("價格 > 50 日均線", bool(last["Adj Close"] > last["SMA_50_Adj"])),
             ("50 日均線 > 200 日均線", bool(last["SMA_50_Adj"] > last["SMA_200_Adj"])),
             ("MACD 柱狀圖 > 0", bool(last["MACD_Hist"] > 0)),
             ("RSI14 介乎 40–70", bool(40 <= last["RSI14"] <= 70)),
             ("ATR 數據有效", bool(pd.notna(last["ATR14"])))]
    if s_now == 1 and s_prev == 0:
        state = "符合入場條件（新訊號）"
    elif s_now == 1:
        state = "持有（趨勢條件維持）"
    elif s_prev == 1:
        state = "符合出場條件：" + (sg["Event"].iloc[-1] or "")
    else:
        state = "觀望（未符合入場條件）"
    close, atr = float(last["Close"]), float(last["ATR14"]) if pd.notna(last["ATR14"]) else np.nan
    stop = close - params["stop_mult"] * atr
    target = close + params["target_mult"] * atr
    risk_ps = close - stop
    h60 = float(last["High60"]) if pd.notna(last["High60"]) else np.nan
    rr_atr = (target - close) / risk_ps if risk_ps > 0 else np.nan
    rr_res = (h60 - close) / risk_ps if risk_ps > 0 and pd.notna(h60) and h60 > close else np.nan
    shares = int(np.floor(params["account"] * params["risk_pct"] / 100 / risk_ps)) if risk_ps > 0 else 0
    return {"狀態": state, "日期": d.index[-1].date(), "收市價": close, "ATR14": atr, "止損價": stop,
            "目標價（ATR）": target, "60日高位": h60, "風險回報比（ATR目標）": rr_atr,
            "風險回報比（60日高位）": rr_res, "建議最大股數": shares, "單筆風險金額": shares * risk_ps,
            "倉位市值": shares * close, "條件": conds}


def analyze(df, params):
    d = add_indicators(df)
    sg = generate_signals(d, params["stop_mult"])
    daily, trades, metrics = backtest(d, sg["Signal"], params["cost_bps"])
    return {"data": d.join(sg).join(daily), "trades": trades, "metrics": metrics, "rec": recommend(d, sg, params)}


def build_excel(results, params):
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Font, PatternFill, Alignment
    buf = io.BytesIO()
    summ, ind, met, trd, daily = [], [], [], [], []
    for s, r in results.items():
        rec, d, m = r["rec"], r["data"], r["metrics"].set_index("指標")
        summ.append({"股票": s, "資料日期": rec["日期"], "收市價": rec["收市價"], "規則狀態": rec["狀態"],
                     "止損價": rec["止損價"], "目標價（ATR）": rec["目標價（ATR）"], "60日高位": rec["60日高位"],
                     "風險回報比（ATR目標）": rec["風險回報比（ATR目標）"], "風險回報比（60日高位）": rec["風險回報比（60日高位）"],
                     "建議最大股數": rec["建議最大股數"], "單筆風險金額": rec["單筆風險金額"], "倉位市值": rec["倉位市值"],
                     "策略總回報": m.loc["總回報", "策略"], "買入持有總回報": m.loc["總回報", "買入持有"],
                     "策略最大回撤": m.loc["最大回撤", "策略"], "買入持有最大回撤": m.loc["最大回撤", "買入持有"]})
        l = d.iloc[-1]
        ind.append({"股票": s, "日期": d.index[-1].date(), "收市價": l["Close"], "調整收市價": l["Adj Close"],
                    "SMA20": l["SMA_20_Adj"], "SMA50": l["SMA_50_Adj"], "SMA200": l["SMA_200_Adj"],
                    "MACD": l["MACD"], "MACD Signal": l["MACD_Signal"], "MACD 柱": l["MACD_Hist"], "RSI14": l["RSI14"],
                    "ATR14": l["ATR14"], "布林上軌": l["BB_Upper"], "布林下軌": l["BB_Lower"]})
        mm = r["metrics"].copy(); mm.insert(0, "股票", s); met.append(mm)
        tt = r["trades"].copy(); tt.insert(0, "股票", s); trd.append(tt)
        dd = d.reset_index(); dd.insert(0, "Ticker", s); daily.append(dd)
    eq = pd.concat({**{f"{s} 策略": r["data"]["Strategy_Equity"] for s, r in results.items()},
                    **{f"{s} 買入持有": r["data"]["BuyHold_Equity"] for s, r in results.items()}}, axis=1).sort_index()
    eq.index = eq.index.date
    param_df = pd.DataFrame([("報告生成時間（UTC）", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")),
                             ("止損倍數 ×ATR", params["stop_mult"]), ("目標倍數 ×ATR", params["target_mult"]),
                             ("單邊交易成本（bps）", params["cost_bps"]), ("帳戶資金（USD）", params["account"]),
                             ("單筆風險（%）", params["risk_pct"]), ("數據來源", "Yahoo Finance（yfinance），日線")] +
                            [(a, b) for a, b in RULES], columns=["項目", "內容"])
    disc = pd.DataFrame({"免責聲明": ["本報告由規則化程式自動生成，僅供研究參考，不構成投資建議或招攬。",
                                    "回測受樣本區間、數據品質及參數選擇影響，結果可能過度擬合，且不代表未來表現。",
                                    "未包含滑點、稅項、匯率、融資成本與流動性風險；請自行核對數據及獨立判斷。"]})
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        pd.DataFrame(summ).to_excel(w, sheet_name="摘要與建議", index=False)
        param_df.to_excel(w, sheet_name="交易策略", index=False)
        pd.DataFrame(ind).to_excel(w, sheet_name="技術指標", index=False)
        pd.concat(met).to_excel(w, sheet_name="歷史績效", index=False)
        pd.concat(trd).to_excel(w, sheet_name="交易紀錄", index=False)
        eq.to_excel(w, sheet_name="權益曲線", index_label="日期")
        pd.concat(daily).to_excel(w, sheet_name="每日數據", index=False)
        disc.to_excel(w, sheet_name="免責聲明", index=False)
        head = PatternFill("solid", fgColor="1F3864")
        pct_heads = {"策略總回報", "買入持有總回報", "策略最大回撤", "買入持有最大回撤", "扣成本後回報", "Return_1D",
                     "Drawdown_From_Observed_Peak", "Strategy_Return"}
        for ws in w.book.worksheets:
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font, c.fill = Font(bold=True, color="FFFFFF"), head
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
            heads = {c.value: c.column for c in ws[1]}
            for h in pct_heads & set(heads):
                for row in range(2, ws.max_row + 1):
                    ws.cell(row, heads[h]).number_format = "0.00%"
            for h, ci in heads.items():
                if h in ("收市價", "止損價", "目標價（ATR）", "60日高位", "風險回報比（ATR目標）", "風險回報比（60日高位）", "單筆風險金額", "倉位市值"):
                    for row in range(2, ws.max_row + 1):
                        ws.cell(row, ci).number_format = "#,##0.00"
        wm = w.book["歷史績效"]
        for row in range(2, wm.max_row + 1):
            if wm.cell(row, 2).value in PCT_METRICS:
                for col in (3, 4):
                    wm.cell(row, col).number_format = "0.00%"
            else:
                for col in (3, 4):
                    wm.cell(row, col).number_format = "0.00"
        wt = w.book["交易策略"]
        wt.column_dimensions["B"].width = 110
        for row in wt.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        we = w.book["權益曲線"]
        if we.max_row > 2:
            ch = LineChart(); ch.title = "策略 vs 買入持有（起點 = 1）"; ch.height, ch.width = 12, 26
            ch.add_data(Reference(we, min_col=2, max_col=min(we.max_column, 17), min_row=1, max_row=we.max_row), titles_from_data=True)
            ch.set_categories(Reference(we, min_col=1, min_row=2, max_row=we.max_row))
            we.add_chart(ch, "H2")
    return buf.getvalue()
